"""Sentetik örnek: 125x100 çerçeve tanımı (xlsx) + buna uygun binary dosya. Gerçek veri içermez."""
import random, struct, math, sys, os
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter

out = sys.argv[1] if len(sys.argv) > 1 else 'ornek'
os.makedirs(out, exist_ok=True)
R, C = 125, 100; L = R * C
rnd = random.Random(7)
names = ['BASLIK', 'GUC', 'ITKI', 'NAVIG', 'HABER', 'ISIL', 'YAKIT', 'KONTROL', 'SENSOR', 'YUK', 'KAMERA', 'DEPO']
# 12 alt sistem, rastgele bayt sayıları (her biri >= 150, toplam 12500)
MIN = 150
rest = L - MIN * len(names)
cuts = sorted(rnd.randint(0, rest) for _ in range(len(names) - 1))
parts = [b - a for a, b in zip([0] + cuts, cuts + [rest])]
sizes = [MIN + p for p in parts]
assert sum(sizes) == L
cell = []
for nm, sz in zip(names, sizes): cell += [nm] * sz
start = {}; o = 0
for nm, sz in zip(names, sizes): start[nm] = o; o += sz
print('Alt sistem bayt sayıları:', dict(zip(names, sizes)))

pal = ['FFE699', 'C6E0B4', 'BDD7EE', 'F8CBAD', 'D9D2E9', 'B4E5E0', 'FFF2CC', 'E2EFDA', 'DDEBF7', 'FCE4D6', 'EAD1DC', 'D0E0E3']
wb = Workbook(); ws = wb.active; ws.title = 'Çerçeve'
fills = {n: PatternFill('solid', fgColor=pal[i]) for i, n in enumerate(names)}
for r in range(R):
    for c in range(C):
        nm = cell[r * C + c]; x = ws.cell(r + 1, c + 1, nm); x.fill = fills[nm]; x.font = Font(size=8); x.alignment = Alignment(horizontal='center')
for c in range(C): ws.column_dimensions[get_column_letter(c + 1)].width = 9
ws.freeze_panes = 'A1'

sc = wb.create_sheet('Senaryo')
for row in [('Anahtar', 'Değer'), ('Okuma sırası', 'satır'), ('Senkron (hex)', 'EB 90 A5 C3'), ('Kare hızı (Hz)', 50), ('Atlanacak bayt', 0), ('Sayaç alanı', 'SAYAC')]: sc.append(row)
sc.column_dimensions['A'].width = 22; sc.column_dimensions['B'].width = 20

al = wb.create_sheet('Alanlar')
al.append(('Ad', 'Başlangıç', 'Tür', 'Bayt sırası', 'Bit', 'İfade', 'Birim'))
FIELDS = [
 ('SAYAC', 'BASLIK_5', 'u16', 'BE', '', '', ''),
 ('GUC_VOLT', 'GUC_1', 'u16', 'BE', '', 'x*0.001', 'V'),
 ('GUC_AKIM', 'GUC_3', 'i16', 'BE', '', 'x*0.01', 'A'),
 ('ISIL_T1', 'ISIL_1', 'i16', 'LE', '', 'x*0.1-40', 'C'),
 ('NAVIG_ENLEM', 'NAVIG_1', 'f32', 'BE', '', '', 'derece'),
 ('YAKIT_SEVIYE', 'YAKIT_1', 'u16', 'BE', '', 'x/655.35', '%'),
 ('KONTROL_MOD', 'KONTROL_1', 'u8', 'BE', '4-7', '', ''),
 ('CBIT_0', 'KONTROL_1', 'u8', 'BE', '0', '', ''),
 ('CBIT_1', 'KONTROL_1', 'u8', 'BE', '1', '', ''),
 ('CBIT_2', 'KONTROL_1', 'u8', 'BE', '2', '', ''),
 ('GUC_WATT', 'GUC_1', 'u16', 'BE', '', 'GUC_VOLT*GUC_AKIM', 'W'),
]
for f in FIELDS: al.append(f)
for col, w in zip('ABCDEFG', (16, 14, 8, 12, 8, 24, 8)): al.dimensions
wb.save(f'{out}/ornek_cerceve.xlsx')

# --- binary ---
SYNC = bytes([0xEB, 0x90, 0xA5, 0xC3]); NF = 200; LOST = {80, 81, 150}  # bu kareler dosyaya yazılmaz
buf = bytearray()
for k in range(NF):
    fr = bytearray(rnd.randrange(256) for _ in range(L))
    fr[0:4] = SYNC
    def put(nm, idx, fmt, v): o = start[nm] + idx - 1; fr[o:o + struct.calcsize(fmt)] = struct.pack(fmt, v)
    t = k / 50
    put('BASLIK', 5, '>H', k)
    put('GUC', 1, '>H', int(28000 + 800 * math.sin(2 * math.pi * 0.5 * t) + rnd.gauss(0, 15)))
    put('GUC', 3, '>h', int(1200 + 300 * math.sin(2 * math.pi * 0.2 * t)))
    put('ISIL', 1, '<h', int((25 + 0.05 * k - -40 + 5 * math.sin(t)) * 10))
    put('NAVIG', 1, '>f', 39.9 + 0.001 * k)
    put('YAKIT', 1, '>H', int((100 - 0.2 * k) * 655.35))
    fr[start['KONTROL']] = ((k // 20 % 4) << 4) | (1 if k % 10 < 5 else 0) | (2 if k % 20 < 10 else 0) | (4 if k > 120 else 0)
    if k == 60: put('GUC', 1, '>H', 65000)          # aykırı nokta
    if k in LOST: continue
    buf += fr
open(f'{out}/ornek_veri.bin', 'wb').write(buf)
print('binary:', len(buf), 'bayt,', len(buf) // L, 'kare')
